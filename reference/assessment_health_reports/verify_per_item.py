"""
Checks a generated Per-Item Speededness/Omission/Completion workbook against
per_item_export_spec.json (extracted from the reference workbook).

Data-agnostic: row counts, assessment names and values are allowed to differ
from the reference; STRUCTURE, STYLE, SIZES, TABLES and CONDITIONAL FORMATTING
must match. Also enforces the lessons from earlier exports:
  - no width-less <col> stubs (render as hidden columns in some clients)
  - no hidden columns / rows
  - explicit row heights + sheetFormatPr defaultRowHeight
  - no placeholder strings ("Not sourced", "n/a", ...) in cells
  - conditional-formatting + tableParts element order inside worksheet XML

Usage:
    python3 verify_per_item.py <generated.xlsx> [per_item_export_spec.json]
Exit code 0 = PASS, 1 = FAIL.
"""
import sys, re, json, math, zipfile
import openpyxl
from openpyxl.utils import get_column_letter as L

PLACEHOLDER = re.compile(r"not sourced|not available|^n/?a$|^tbd$|^unknown$|^#", re.I)
CT_ORDER = ["sheetPr", "dimension", "sheetViews", "sheetFormatPr", "cols", "sheetData",
            "sheetCalcPr", "sheetProtection", "protectedRanges", "scenarios", "autoFilter",
            "sortState", "dataConsolidate", "customSheetViews", "mergeCells", "phoneticPr",
            "conditionalFormatting", "dataValidations", "hyperlinks", "printOptions",
            "pageMargins", "pageSetup", "headerFooter", "rowBreaks", "colBreaks",
            "customProperties", "cellWatches", "ignoredErrors", "smartTags", "drawing",
            "legacyDrawing", "legacyDrawingHF", "picture", "oleObjects", "controls",
            "webPublishItems", "tableParts", "extLst"]


def sheet_xml_map(zf):
    wb = zf.read("xl/workbook.xml").decode()
    rels = zf.read("xl/_rels/workbook.xml.rels").decode()
    rid2target = {}
    for m in re.finditer(r"<Relationship\b[^>]*>", rels):
        t = m.group(0)
        i = re.search(r'Id="([^"]+)"', t); g = re.search(r'Target="([^"]+)"', t)
        if i and g:
            rid2target[i.group(1)] = g.group(1).lstrip("/").replace("xl/", "", 1)
    out = {}
    for m in re.finditer(r"<sheet\b[^>]*>", wb):
        t = m.group(0)
        n = re.search(r'name="([^"]+)"', t); r = re.search(r'r:id="([^"]+)"', t)
        if n and r:
            name = n.group(1).replace("&amp;", "&")
            out[name] = "xl/" + rid2target[r.group(1)]
    return out


def raw_cols(xml):
    m = re.search(r"<cols>(.*?)</cols>", xml, re.S)
    res = {}
    if not m:
        return res
    for tag in re.findall(r"<col\b[^>]*/>", m.group(1)):
        mn = int(re.search(r'min="(\d+)"', tag).group(1)); mx = int(re.search(r'max="(\d+)"', tag).group(1))
        w = re.search(r'width="([\d.]+)"', tag)
        hidden = re.search(r'hidden="(1|true)"', tag) is not None
        for c in range(mn, mx + 1):
            res[L(c)] = {"width": float(w.group(1)) if w else None, "hidden": hidden}
    return res


def top_level_order(xml):
    body = re.search(r"<worksheet\b[^>]*>(.*)</worksheet>", xml, re.S).group(1)
    names, depth, i = [], 0, 0
    for m in re.finditer(r"<(/?)([A-Za-z0-9:]+)([^>]*?)(/?)>", body):
        close, name, _, selfc = m.groups()
        if close:
            depth -= 1
        else:
            if depth == 0:
                names.append(name.split(":")[-1])
            if not selfc:
                depth += 1
    return names


def check(path, spec_path):
    spec = json.load(open(spec_path))
    prob = []
    F = lambda m: prob.append(m)
    wb = openpyxl.load_workbook(path)
    zf = zipfile.ZipFile(path)
    smap = sheet_xml_map(zf)
    item = spec["item_sheet"]; rd = spec["readme"]

    if wb.sheetnames[0] != "README & Methodology":
        F(f"first sheet must be 'README & Methodology', got {wb.sheetnames[0]!r}")
    if len(wb.sheetnames) < 2:
        F("no item sheets found")

    for name in wb.sheetnames:
        ws = wb[name]; xml = zf.read(smap[name]).decode("utf-8")
        tag = f"[{name}]"
        # --- generic ---
        if not re.search(r"<sheetFormatPr[^>]*defaultRowHeight=\"13\.8\"", xml):
            F(f"{tag} sheetFormatPr defaultRowHeight=13.8 missing")
        if ws.freeze_panes: F(f"{tag} freeze panes must be none (got {ws.freeze_panes})")
        if ws.sheet_properties.tabColor is not None: F(f"{tag} tab color must be none")
        cols = raw_cols(xml)
        for c, d in cols.items():
            if d["hidden"]: F(f"{tag} column {c} is hidden")
            if d["width"] is None: F(f"{tag} column {c} is a width-less <col> STUB (renders hidden in some clients)")
        for r, rd_ in ws.row_dimensions.items():
            if rd_.hidden: F(f"{tag} row {r} is hidden")
        order = top_level_order(xml)
        idx = [CT_ORDER.index(n) for n in order if n in CT_ORDER]
        if idx != sorted(idx): F(f"{tag} worksheet child elements out of schema order: {order}")
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and PLACEHOLDER.search(c.value.strip()):
                    F(f"{tag} placeholder text {c.coordinate}={c.value!r}")

        if name == "README & Methodology":
            if sorted(str(r) for r in ws.merged_cells.ranges) != sorted(rd["merges"]): F(f"{tag} merges differ: {sorted(str(r) for r in ws.merged_cells.ranges)}")
            for c, w in rd["col_widths"].items():
                got = cols.get(c, {}).get("width")
                if got is None or abs(got - w) > 0.01: F(f"{tag} width {c}: expected {w}, got {got}")
            for r, h in rd["row_heights"].items():
                got = ws.row_dimensions[int(r)].height
                if got is None or abs(got - h) > 0.01: F(f"{tag} row {r} height: expected {h}, got {got}")
            if ws["A1"].value != rd["title"]: F(f"{tag} title text differs")
            if [ws.cell(row=4, column=c).value for c in range(1, 5)] != rd["table_header"]: F(f"{tag} table header differs")
            for i, row in enumerate(rd["rows"]):
                got = [ws.cell(row=5 + i, column=c).value for c in range(1, 5)]
                if got != row: F(f"{tag} README row {5+i} text differs")
            if ws["A1"].fill.fgColor.rgb != "FFB2375B" or ws["A1"].font.name != "Carlito": F(f"{tag} title style differs")
            continue

        # --- item sheets ---
        n_rows = ws.max_row - 4
        last = ws.max_row
        if sorted(str(r) for r in ws.merged_cells.ranges) != sorted(item["merges"]): F(f"{tag} merges differ: {sorted(str(r) for r in ws.merged_cells.ranges)}")
        exp_title = item["title_template"].format(AssessmentName=name)
        if ws["A1"].value != exp_title: F(f"{tag} A1 title: expected {exp_title!r}, got {ws['A1'].value!r}")
        if ws["A2"].value != item["subtitle"]: F(f"{tag} A2 subtitle differs")
        if ws["A1"].fill.fgColor.rgb != "FFB2375B" or ws["A1"].font.color.rgb != "FFFFFFFF": F(f"{tag} title fill/font differ")
        if ws["A2"].fill.fgColor.rgb != "FFF9F5F2": F(f"{tag} subtitle fill differs")
        for c in item["columns"]:
            if ws[f"{c['col']}4"].value != c["header"]: F(f"{tag} header {c['col']}4: expected {c['header']!r}, got {ws[c['col']+'4'].value!r}")
        # widths
        for c, w in item["col_widths_fixed"].items():
            got = cols.get(c, {}).get("width")
            if got is None or abs(got - w) > 0.01: F(f"{tag} width {c}: expected {w}, got {got}")
        for c, v in item["col_widths_variable"].items():
            got = cols.get(c, {}).get("width")
            if got is None or got + 0.01 < v["min"] or got > v["max"] + 0.01:
                F(f"{tag} width {c}: expected between {v['min']} and {v['max']}, got {got}")
        # heights
        rh = item["row_heights"]
        for r in (1, 2, 3, 4):
            exp = rh[str(r)]; got = ws.row_dimensions[r].height
            tol = 2.0 if r == 1 else 0.01  # title row varies 33-35.4 in the reference
            if got is None or abs(got - exp) > tol: F(f"{tag} row {r} height: expected ~{exp}, got {got}")
        for r in range(5, last + 1):
            got = ws.row_dimensions[r].height
            if got is None or abs(got - rh["data_rows"]) > 0.01:
                F(f"{tag} data row {r} height: expected {rh['data_rows']}, got {got}"); break
        # table
        tbls = list(ws.tables.values())
        if len(tbls) != 1: F(f"{tag} expected exactly 1 Excel Table, found {len(tbls)}")
        else:
            t = tbls[0]
            exp_name = re.sub(r"[^A-Za-z0-9]", "", name)[:20] + "ItemTable"
            if t.name != exp_name: F(f"{tag} table name: expected {exp_name}, got {t.name}")
            if t.ref != f"A4:T{last}": F(f"{tag} table ref: expected A4:T{last}, got {t.ref}")
            if [c.name for c in t.tableColumns] != [c["header"] for c in item["columns"]]: F(f"{tag} table columns differ from headers")
        # CF
        cf = {str(k.sqref): v for k, v in ws.conditional_formatting._cf_rules.items()}
        for spec_cf in item["conditional_formatting"]:
            rng = spec_cf["range"].replace("<last>", str(last))
            if rng not in cf: F(f"{tag} CF range {rng} missing (found {list(cf)})"); continue
            exp_f = [r[0] for r in spec_cf["rules"]]
            if [r.formula[0] for r in cf[rng]] != exp_f: F(f"{tag} CF formulas on {rng} differ: {[r.formula for r in cf[rng]]}")
            for r, (_, tier) in zip(cf[rng], spec_cf["rules"]):
                d = wb._differential_styles.styles[r.dxfId]
                col = (d.fill.bgColor.rgb if d.fill.bgColor.rgb not in (None, "00000000") else d.fill.fgColor.rgb)
                if col != item["cf_colors"][tier]["fill"]: F(f"{tag} CF {rng} {tier} fill {col} != {item['cf_colors'][tier]['fill']}")
                if (d.font.color.rgb if d.font and d.font.color else None) != item["cf_colors"][tier]["font"]: F(f"{tag} CF {rng} {tier} font color differs")
        if set(cf) - {s["range"].replace("<last>", str(last)) for s in item["conditional_formatting"]}:
            F(f"{tag} unexpected extra CF ranges: {list(cf)}")
        # cell formats + data rules
        late_n = math.ceil(0.25 * n_rows)
        prev_qpn = None
        for r in range(5, last + 1):
            for c in item["columns"]:
                cell = ws[f"{c['col']}{r}"]
                if cell.font.name != "Carlito" or cell.font.size != 11: F(f"{tag} {cell.coordinate} font must be Carlito 11"); break
                if cell.number_format != c["number_format"]: F(f"{tag} {cell.coordinate} number_format {cell.number_format!r} != {c['number_format']!r}"); break
                a = c["data_align"]
                if (cell.alignment.horizontal, cell.alignment.vertical, bool(cell.alignment.wrap_text)) != (a["h"], a["v"], bool(a["wrap"])):
                    F(f"{tag} {cell.coordinate} alignment differs"); break
                if cell.value is None: F(f"{tag} {cell.coordinate} is blank")
            g = lambda col: ws[f"{col}{r}"].value
            if g("A") != name: F(f"{tag} A{r} AssessmentName != sheet name")
            sec_late = (r - 5) >= n_rows - late_n
            if g("I") != item["test_section_labels"][1 if sec_late else 0]: F(f"{tag} I{r} Test Section label wrong for position (late = last {late_n} rows)")
            q = g("H")
            if isinstance(q, (int, float)):
                if prev_qpn is not None and q < prev_qpn: F(f"{tag} H{r} rows not sorted by QuestionPresentedNumber")
                prev_qpn = q
            s = g("N"); sl = item["status_labels"]
            if isinstance(s, (int, float)):
                exp = sl["speededness"][0] if s <= 0.05 else sl["speededness"][1] if s <= 0.15 else sl["speededness"][2]
                if g("O") != exp: F(f"{tag} O{r} speededness status {g('O')!r} inconsistent with N={s}")
                tier = ["Good", "Review", "Flag"][sl["speededness"].index(exp)]
                exp_note = item["notes_templates"]["late" if sec_late else "early_middle"][tier]
                if g("T") != exp_note: F(f"{tag} T{r} note does not match template for (section, status)")
            p = g("P")
            if isinstance(p, (int, float)):
                exp = sl["omission"][0] if p <= 0.05 else sl["omission"][1] if p <= 0.10 else sl["omission"][2]
                if g("R") != exp: F(f"{tag} R{r} omission status inconsistent with P={p}")
            qv = g("Q")
            if isinstance(qv, (int, float)):
                exp = sl["completion"][0] if qv >= 0.95 else sl["completion"][1] if qv >= 0.90 else sl["completion"][2]
                if g("S") != exp: F(f"{tag} S{r} completion status inconsistent with Q={qv}")
                if isinstance(p, (int, float)) and abs((p + qv) - 1) > 1e-6: F(f"{tag} row {r}: Omission + Completion != 1")
    return prob


if __name__ == "__main__":
    path = sys.argv[1]
    spec = sys.argv[2] if len(sys.argv) > 2 else "per_item_export_spec.json"
    problems = check(path, spec)
    if problems:
        print(f"FAIL -- {len(problems)} issue(s) in {path}")
        for p in problems[:200]:
            print(" -", p)
        sys.exit(1)
    print(f"PASS -- {path} matches the per-item structure/style spec")
